"""Build the current progress report from its Markdown source."""

from docx import Document
from docx.shared import Pt

from create_handover_docx import ROOT, add_inline, configure_document, parse_markdown, set_run_font


SOURCE = ROOT / "docs" / "BAO-CAO-TINH-TRANG-2026-10-07.md"
OUTPUT = ROOT / "docs" / "Bao-cao-danh-gia-tien-do-May-VSTEP.docx"


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.clear()
    set_run_font(header.add_run("MÂY VSTEP  |  BÁO CÁO TÌNH TRẠNG"), size=7.8, bold=True)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("Báo cáo tình trạng Mây VSTEP"), size=24, bold=True, color="000000")

    date = doc.add_paragraph()
    date.paragraph_format.space_after = Pt(14)
    set_run_font(date.add_run("Cập nhật ngày 07 tháng 10 năm 2026"), size=10)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    add_inline(intro, lines[2], 10.2, "211D20")
    parse_markdown(doc, lines)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("3. Bằng chứng kiểm tra"):
            paragraph.paragraph_format.page_break_before = True
            break

    doc.core_properties.title = "Báo cáo tình trạng Mây VSTEP"
    doc.core_properties.subject = "Bản năm đề nhập, kiểm thử và phát hành"
    doc.core_properties.author = "Dự án Mây VSTEP"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
