from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "HANDOVER.md"
OUTPUT = ROOT / "docs" / "Bao-cao-ban-giao-May-VSTEP.docx"

INK = "211D20"
MUTED = "6E6268"
ROSE = "C93B72"
BLUSH = "FCEEF4"
PALE = "FAF7F8"
BORDER = "D9D9D9"
DARK = "4A3B42"


def set_run_font(run, name="Arial", size=None, bold=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(element, fill):
    props = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=120, bottom=85, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, display, end])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ROSE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|https?://\S+)")


def add_inline(paragraph, text, default_size=10.2, color=INK):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=default_size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", default_size - 0.5, color=DARK)
            run.font.highlight_color = None
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            add_hyperlink(paragraph, token.rstrip(".,"), token.rstrip(".,"))
            suffix = token[len(token.rstrip(".,")) :]
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, size=default_size, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=default_size, color=color)


def add_markdown_table(doc, rows):
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    usable = 6.75
    if column_count == 3:
        widths = [usable * 0.23, usable * 0.57, usable * 0.20]
    elif column_count == 4:
        widths = [usable * 0.19, usable * 0.18, usable * 0.23, usable * 0.40]
    elif column_count == 2:
        widths = [usable * 0.31, usable * 0.69]
    else:
        widths = [usable / column_count] * column_count
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        prevent_row_split(row)
        if i == 0:
            repeat_header(row)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if i == 0:
                shade(cell, DARK)
            elif i % 2 == 0:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if j >= len(row_data):
                continue
            add_inline(
                p,
                row_data[j],
                default_size=8.1,
                color="FFFFFF" if i == 0 else INK,
            )
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_architecture(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, "Luồng chính của dữ liệu học tập", default_size=9.5, color=MUTED)
    headers = ["Tầng giao diện", "Tầng trạng thái", "Lưu trữ chính"]
    labels = ["Next.js và React UI", "StudyProvider và store", "localStorage cho state học tập"]
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    widths = [2.05, 2.20, 2.45]
    for i, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 120, 80, 120, 80)
        shade(cell, DARK)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_run_font(run, size=8.3, bold=True, color="FFFFFF")
    for i, (cell, label) in enumerate(zip(table.rows[1].cells, labels)):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 130, 90, 130, 90)
        shade(cell, BLUSH)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_run_font(run, size=8.7, bold=True, color=ROSE)
    set_table_borders(table, color="E8CDD8")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    add_inline(
        p,
        "Nhánh riêng: giao diện → IndexedDB cho audio | store → Zod → Supabase Auth/RPC → study_snapshots có RLS",
        default_size=8.8,
        color=MUTED,
    )


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("MÂY VSTEP")
    set_run_font(run, size=10, bold=True, color=ROSE)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Báo cáo bàn giao website Mây VSTEP")
    set_run_font(run, size=29, bold=True, color="000000")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("Trạng thái sản phẩm, kiến trúc, vận hành và lộ trình phát triển")
    set_run_font(run, size=13, color=MUTED)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table, color="E6DDE1")
    header_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    repeat_header(table.rows[0])
    shade(header_cell, DARK)
    header_p = header_cell.paragraphs[0]
    header_p.paragraph_format.space_after = Pt(0)
    header_run = header_p.add_run("THÔNG TIN BÀN GIAO")
    set_run_font(header_run, size=8.5, bold=True, color="FFFFFF")
    metadata = [
        ("Ngày chốt", "10/09/2026"),
        ("Mốc đánh giá", "9f4d8ea + vòng cải tiến learning loop 10/09/2026"),
        ("Nhánh", "main"),
        ("Repository", "https://github.com/binhphanbp/vstep"),
        ("Người học", "Gùa, tên ở nhà Rùa"),
    ]
    for i, (label, value) in enumerate(metadata, start=1):
        left, right = table.rows[i].cells
        left.width = Inches(1.35)
        right.width = Inches(5.1)
        shade(left, BLUSH)
        if i % 2:
            shade(right, PALE)
        for cell in (left, right):
            set_cell_margins(cell, 130, 150, 130, 150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        left_p = left.paragraphs[0]
        left_p.paragraph_format.space_after = Pt(0)
        add_inline(left_p, label, 9, ROSE)
        for run in left_p.runs:
            run.bold = True
        right_p = right.paragraphs[0]
        right_p.paragraph_format.space_after = Pt(0)
        if value.startswith("http"):
            add_hyperlink(right_p, "github.com/binhphanbp/vstep", value)
        else:
            add_inline(right_p, value, 9, INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Trạng thái tại thời điểm bàn giao")
    set_run_font(run, size=10, bold=True, color=ROSE)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    add_inline(
        p,
        "Bản code hiện đủ để Gùa pilot hằng ngày trên local, Supabase thật đã kết nối và kiểm thử, toàn bộ pipeline kiểm tra đang đạt. Daily Mission đã giải thích lý do chọn bài và ưu tiên lỗi đến hạn hoặc sai với mức tự tin cao. Các điều kiện còn thiếu để gọi là production hoàn chỉnh là URL HTTPS, nghiệm thu thiết bị thật, thử đồng bộ hai thiết bị và thẩm định học liệu bởi giáo viên VSTEP.",
        11,
        INK,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run("Tài liệu bàn giao nội bộ • Không chứa mật khẩu hoặc secret key")
    set_run_font(run, size=8.5, color=MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Nội dung bàn giao")
    sections = [
        "Kết luận và hướng đi sản phẩm",
        "Nỗi đau người học và phạm vi chức năng",
        "UX UI và nhận diện",
        "Kiến trúc kỹ thuật và mô hình dữ liệu",
        "Supabase và bảo mật",
        "Kiểm thử và bằng chứng chất lượng",
        "Giới hạn hiện tại",
        "Hướng dẫn vận hành và sao lưu",
        "Roadmap sau bàn giao",
        "Cấu trúc mã nguồn, quyết định và checklist quyền truy cập",
    ]
    for item in sections:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, item, 10.2, INK)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, before, after in (
        ("Heading 1", 17, 15, 7),
        ("Heading 2", 12.5, 11, 5),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_p_pr = title_style.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Cm(0.55)
        style.paragraph_format.first_line_indent = Cm(-0.3)
        style.paragraph_format.space_after = Pt(4)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("MÂY VSTEP  |  BÁO CÁO BÀN GIAO")
    set_run_font(run, size=7.8, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    doc.core_properties.title = "Báo cáo bàn giao website Mây VSTEP"
    doc.core_properties.subject = "Trạng thái sản phẩm, kiến trúc, vận hành và lộ trình phát triển"
    doc.core_properties.author = "Dự án Mây VSTEP"
    doc.core_properties.keywords = "VSTEP, bàn giao, Next.js, Supabase, Gùa"


def parse_markdown(doc, lines):
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    i = start
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                if code_lines and code_lines[0].startswith("flowchart"):
                    add_architecture(doc)
                else:
                    for code_line in code_lines:
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent = Cm(0.45)
                        p.paragraph_format.right_indent = Cm(0.25)
                        p.paragraph_format.space_after = Pt(0)
                        shade(p._p.get_or_add_pPr(), PALE)
                        run = p.add_run(code_line or " ")
                        set_run_font(run, "Consolas", 8.6, color=DARK)
                    doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:]
            if title.startswith("17. "):
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, title, 17, "000000")
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], 12.5, "000000")
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for row_line in table_lines:
                cells = [c.strip() for c in row_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            add_markdown_table(doc, rows)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.62)
            p.paragraph_format.first_line_indent = Cm(-0.42)
            p.paragraph_format.space_after = Pt(4)
            prefix = p.add_run(f"{numbered.group(1)}. ")
            set_run_font(prefix, size=10.2, color=INK)
            add_inline(p, numbered.group(2), 10.2, INK)
        elif bullet:
            content = bullet.group(1)
            if content.startswith("[x]"):
                content = "[x] " + content[3:].lstrip()
            elif content.startswith("[ ]"):
                content = "[ ] " + content[3:].lstrip()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, content, 10.2, INK)
        else:
            p = doc.add_paragraph()
            add_inline(p, line, 10.2, INK)
        i += 1


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8").splitlines())
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
